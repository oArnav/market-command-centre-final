import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
import pandas as pd
import yfinance as yf

WATCHLIST_CSV = os.path.join(os.path.dirname(__file__), "watchlist.csv")
OUTPUT_XLSX = os.path.join(os.path.dirname(__file__), "stock_baseline_matrix.xlsx")

TICKER_ALIAS_MAP = {
    "TATAMOTORS.NS": ["TMPV.NS", "TATAMOTORS.NS"],
    "TATAMOTORS": ["TMPV.NS", "TATAMOTORS.NS"],
    "ZOMATO.NS": ["ETERNAL.NS", "ZOMATO.NS"],
    "ZOMATO": ["ETERNAL.NS", "ZOMATO.NS"]
}

def fetch_quick_price(symbol):
    clean = str(symbol).strip().upper()
    formatted = clean if clean.endswith(".NS") else f"{clean}.NS"
    aliases = TICKER_ALIAS_MAP.get(clean, TICKER_ALIAS_MAP.get(formatted, [formatted, clean]))
    for sym in aliases:
        try:
            t = yf.Ticker(sym)
            if hasattr(t, "fast_info") and "lastPrice" in t.fast_info:
                p = t.fast_info["lastPrice"]
                if p and not pd.isna(p):
                    return round(float(p), 2)
            hist = t.history(period="1d")
            if not hist.empty and "Close" in hist.columns:
                p = hist["Close"].iloc[-1]
                if p and not pd.isna(p):
                    return round(float(p), 2)
        except Exception:
            continue
    return 1000.00

def create_stock_matrix_from_csv(csv_path=WATCHLIST_CSV, output_path=OUTPUT_XLSX):
    if not os.path.exists(csv_path):
        print(f"Watchlist CSV not found at: {csv_path}")
        return

    df = pd.read_csv(csv_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Watchlist Tracker (Version 1)"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Stock (NSE Ticker)", 
        "Current Price (₹)", 
        "Buy Below (₹)", 
        "Sell Above (₹)", 
        "Status", 
        "Your One-Line Thesis"
    ]
    ws.append(headers)

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    total_rows = len(df)
    for idx, r in df.iterrows():
        row_idx = idx + 2
        sym = str(r["nse_ticker"]).strip().upper()
        buy_val = float(r["buy_below"])
        sell_val = float(r["sell_above"])
        thesis_val = str(r["thesis"])
        curr_price = fetch_quick_price(sym)

        # Exact formula required: =IF(CurrentPrice<=BuyTarget,"BUY ALERT",IF(CurrentPrice>=SellTarget,"SELL ALERT","Watching"))
        status_formula = f'=IF(B{row_idx}<=C{row_idx}, "BUY ALERT", IF(B{row_idx}>=D{row_idx}, "SELL ALERT", "Watching"))'

        row_data = [sym, curr_price, buy_val, sell_val, status_formula, thesis_val]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 32

        # Col A: Ticker
        ws.cell(row=row_idx, column=1).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")

        # Col B, C, D: Currency
        for c in [2, 3, 4]:
            cell = ws.cell(row=row_idx, column=c)
            cell.number_format = '"₹"#,##0.00'
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(horizontal="right", vertical="center")

        # Col E: Status
        ws.cell(row=row_idx, column=5).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=row_idx, column=5).alignment = Alignment(horizontal="center", vertical="center")

        # Col F: Thesis
        ws.cell(row=row_idx, column=6).font = Font(name="Calibri", size=10, italic=True)
        ws.cell(row=row_idx, column=6).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for c in range(1, 7):
            ws.cell(row=row_idx, column=c).border = thin_border

    # Conditional Formatting for Column E (Status)
    if total_rows > 0:
        end_row = total_rows + 1
        range_str = f"E2:E{end_row}"

        buy_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        buy_font = Font(color="006100", bold=True)
        buy_rule = CellIsRule(operator='equal', formula=['"BUY ALERT"'], stopIfTrue=True, fill=buy_fill, font=buy_font)

        sell_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        sell_font = Font(color="9C0006", bold=True)
        sell_rule = CellIsRule(operator='equal', formula=['"SELL ALERT"'], stopIfTrue=True, fill=sell_fill, font=sell_font)

        ws.conditional_formatting.add(range_str, buy_rule)
        ws.conditional_formatting.add(range_str, sell_rule)

    # Column dimensions
    col_widths = {'A': 18, 'B': 20, 'C': 22, 'D': 22, 'E': 20, 'F': 65}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(output_path)
    print(f"Successfully generated dynamic Excel workbook for {total_rows} stocks: {output_path}")

if __name__ == "__main__":
    create_stock_matrix_from_csv()
